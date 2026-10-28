from __future__ import annotations

import html
import io
import json
import math
import os
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from markdown_it import MarkdownIt
from PIL import Image as PILImage

from .model_review_docx import Model3DEmbedError, PLACEHOLDER, embed_models3d


ALLOWED_PREVIEW_COUNTS = {6, 8, 10, 12, 14}
VIEW_PRESETS = [
    ("front", "前视图"),
    ("back", "后视图"),
    ("left", "左视图"),
    ("right", "右视图"),
    ("top", "顶视图"),
    ("bottom", "底视图"),
    ("front_top_right", "右前上等轴视图"),
    ("back_bottom_left", "左后下等轴视图"),
    ("front_top_left", "左前上等轴视图"),
    ("back_bottom_right", "右后下等轴视图"),
    ("front_bottom_right", "右前下等轴视图"),
    ("back_top_left", "左后上等轴视图"),
    ("front_bottom_left", "左前下等轴视图"),
    ("back_top_right", "右后上等轴视图"),
]
MAX_REPORT_TITLE = 160
MAX_IMAGE_BYTES = 12 * 1024 * 1024
MAX_GLB_BYTES = 250 * 1024 * 1024


class ModelReportValidationError(ValueError):
    pass


@dataclass(frozen=True)
class ModelReportRequest:
    path: str
    title: str
    formats: tuple[str, ...]
    preview_count: int
    include_annotations: bool
    annotation_view_count: int
    include_3d: bool
    separate_3d: bool
    separate_models: tuple[tuple[str, float], ...]
    annotation_ids: tuple[str, ...]
    model_max_dimension: float

    @property
    def views(self) -> tuple[tuple[str, str], ...]:
        return tuple(VIEW_PRESETS[: self.preview_count])


def parse_report_request(raw_config: str, annotations: list[dict[str, Any]]) -> ModelReportRequest:
    try:
        payload = json.loads(raw_config)
    except (TypeError, json.JSONDecodeError) as exc:
        raise ModelReportValidationError("config 必须是有效 JSON") from exc
    if not isinstance(payload, dict):
        raise ModelReportValidationError("config 必须是 JSON 对象")

    path = str(payload.get("path") or "").strip()
    title = str(payload.get("title") or "").strip()
    if not title:
        raise ModelReportValidationError("报告标题不能为空")
    if len(title) > MAX_REPORT_TITLE:
        raise ModelReportValidationError(f"报告标题不能超过 {MAX_REPORT_TITLE} 个字符")

    requested_formats = payload.get("formats")
    if not isinstance(requested_formats, list):
        raise ModelReportValidationError("formats 必须是数组")
    formats = tuple(dict.fromkeys(str(item).lower() for item in requested_formats))
    if not formats or any(item not in {"pdf", "docx"} for item in formats):
        raise ModelReportValidationError("formats 只能包含 pdf 或 docx")

    try:
        preview_count = int(payload.get("preview_count"))
    except (TypeError, ValueError) as exc:
        raise ModelReportValidationError("基础预览数量无效") from exc
    if preview_count not in ALLOWED_PREVIEW_COUNTS:
        raise ModelReportValidationError("基础预览数量只能是 6、8、10、12 或 14")

    include_annotations = payload.get("include_annotations") is True
    try:
        annotation_view_count = int(payload.get("annotation_view_count", 1))
    except (TypeError, ValueError) as exc:
        raise ModelReportValidationError("评论视角数量无效") from exc
    if annotation_view_count not in {1, 2, 3, 4}:
        raise ModelReportValidationError("评论视角数量只能是 1、2、3 或 4")
    include_3d = payload.get("include_3d", "docx" in formats) is True
    if include_3d and "docx" not in formats:
        raise ModelReportValidationError("只有 Word 报告可以嵌入交互式 3D 模型")
    separate_3d = payload.get("separate_3d") is True
    raw_separate_models = payload.get("separate_models") or []
    if not isinstance(raw_separate_models, list):
        raise ModelReportValidationError("分别导出的模型列表无效")
    if separate_3d and not include_3d:
        raise ModelReportValidationError("分别导出需要启用 Word 交互式 3D 模型")
    if not separate_3d and raw_separate_models:
        raise ModelReportValidationError("分别导出的模型列表无效")
    if separate_3d and not 2 <= len(raw_separate_models) <= 32:
        raise ModelReportValidationError("分别导出仅适用于 2 至 32 个顶层模型")
    separate_models = []
    for item in raw_separate_models:
        if not isinstance(item, dict):
            raise ModelReportValidationError("分别导出的模型列表无效")
        name = str(item.get("name") or "").strip()
        if not name or len(name) > 100:
            raise ModelReportValidationError("分别导出的模型名称无效")
        try:
            max_dimension = float(item.get("max_dimension"))
        except (TypeError, ValueError) as exc:
            raise ModelReportValidationError("分别导出的模型尺寸无效") from exc
        if not math.isfinite(max_dimension) or max_dimension <= 0:
            raise ModelReportValidationError("分别导出的模型尺寸无效")
        separate_models.append((name, max_dimension))
    requested_annotation_ids = payload.get("annotation_ids") or []
    if not isinstance(requested_annotation_ids, list):
        raise ModelReportValidationError("annotation_ids 必须是数组")
    annotation_ids = tuple(str(item) for item in requested_annotation_ids)
    expected_ids = tuple(item["id"] for item in annotations) if include_annotations else ()
    if annotation_ids != expected_ids:
        raise ModelReportValidationError("评论列表已变化，请重新打开导出设置后再试")

    views = payload.get("views")
    expected_views = [{"id": item[0], "label": item[1]} for item in VIEW_PRESETS[:preview_count]]
    if views != expected_views:
        raise ModelReportValidationError("基础视图清单无效")

    model_max_dimension = 1.0
    if include_3d:
        try:
            model_max_dimension = float(payload.get("model_max_dimension"))
        except (TypeError, ValueError) as exc:
            raise ModelReportValidationError("模型尺寸无效") from exc
        if not math.isfinite(model_max_dimension) or model_max_dimension <= 0:
            raise ModelReportValidationError("模型尺寸无效")

    return ModelReportRequest(
        path=path,
        title=title,
        formats=formats,
        preview_count=preview_count,
        include_annotations=include_annotations,
        annotation_view_count=annotation_view_count,
        include_3d=include_3d,
        separate_3d=separate_3d,
        separate_models=tuple(separate_models),
        annotation_ids=annotation_ids,
        model_max_dimension=model_max_dimension,
    )


def read_validated_image(storage, field_name: str) -> bytes:
    if storage is None:
        raise ModelReportValidationError(f"缺少图片：{field_name}")
    data = storage.read(MAX_IMAGE_BYTES + 1)
    if not data or len(data) > MAX_IMAGE_BYTES:
        raise ModelReportValidationError(f"图片无效或过大：{field_name}")
    try:
        with PILImage.open(io.BytesIO(data)) as image:
            image.verify()
        with PILImage.open(io.BytesIO(data)) as image:
            width, height = image.size
            if width < 320 or height < 180 or width > 4096 or height > 4096:
                raise ModelReportValidationError(f"图片尺寸无效：{field_name}")
            if image.format not in {"JPEG", "PNG"}:
                raise ModelReportValidationError(f"图片格式无效：{field_name}")
    except (OSError, PILImage.DecompressionBombError) as exc:
        raise ModelReportValidationError(f"图片内容无效：{field_name}") from exc
    return data


def read_validated_glb(storage) -> bytes:
    if storage is None:
        raise ModelReportValidationError("Word 导出缺少 GLB 模型")
    data = storage.read(MAX_GLB_BYTES + 1)
    if len(data) > MAX_GLB_BYTES:
        raise ModelReportValidationError("GLB 模型不能超过 250 MB")
    if len(data) < 12 or not data.startswith(b"glTF"):
        raise ModelReportValidationError("上传的模型不是有效 GLB 文件")
    return data


def safe_download_stem(title: str) -> str:
    value = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", "_", title).strip(" ._")
    return value[:100] or "3D审阅报告"


def _markdown_blocks(markdown: str):
    parser = MarkdownIt("commonmark", {"html": False})
    tokens = parser.parse(markdown or "")
    list_stack: list[dict[str, Any]] = []
    in_list_item = False
    heading_level = 0
    for token in tokens:
        if token.type == "heading_open":
            heading_level = int(token.tag[1:]) if token.tag.startswith("h") else 3
        elif token.type == "heading_close":
            heading_level = 0
        elif token.type == "bullet_list_open":
            list_stack.append({"ordered": False, "index": 0})
        elif token.type == "ordered_list_open":
            list_stack.append({"ordered": True, "index": int(token.attrGet("start") or 1) - 1})
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            if list_stack:
                list_stack.pop()
        elif token.type == "list_item_open":
            in_list_item = True
            if list_stack:
                list_stack[-1]["index"] += 1
        elif token.type == "list_item_close":
            in_list_item = False
        elif token.type == "inline":
            if heading_level:
                yield "heading", heading_level, token.children or []
            elif in_list_item and list_stack:
                state = list_stack[-1]
                prefix = f"{state['index']}. " if state["ordered"] else "• "
                yield "list", len(list_stack), [prefix, *(token.children or [])]
            else:
                yield "paragraph", 0, token.children or []
        elif token.type in {"fence", "code_block"}:
            yield "code", 0, token.content


def _inline_html(children: list[Any]) -> str:
    output: list[str] = []
    for token in children:
        if isinstance(token, str):
            output.append(html.escape(token))
        elif token.type == "text":
            output.append(html.escape(token.content))
        elif token.type in {"softbreak", "hardbreak"}:
            output.append("<br/>")
        elif token.type == "strong_open":
            output.append("<b>")
        elif token.type == "strong_close":
            output.append("</b>")
        elif token.type == "em_open":
            output.append("<i>")
        elif token.type == "em_close":
            output.append("</i>")
        elif token.type == "code_inline":
            output.append(f"<font backColor=\"#eeeeee\">{html.escape(token.content)}</font>")
        elif token.type == "link_open":
            href = html.escape(token.attrGet("href") or "", quote=True)
            output.append(f'<link href="{href}" color="#2378a8">')
        elif token.type == "link_close":
            output.append("</link>")
        elif token.type == "image":
            output.append(html.escape(token.content or "[图片略]"))
    return "".join(output)


def _pdf_mixed_text(value: str) -> str:
    parts = re.split(r"([\x00-\x7f]+)", str(value))
    return "".join(
        f'<font name="Helvetica">{html.escape(part)}</font>' if part.isascii() else html.escape(part)
        for part in parts
        if part
    )


def build_pdf(
    request: ModelReportRequest,
    model_filename: str,
    preview_images: dict[str, bytes],
    annotations: list[dict[str, Any]],
    annotation_images: dict[str, list[bytes]],
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = "Helvetica"
    windows_font = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "simhei.ttf"
    if windows_font.is_file():
        try:
            font_name = "ModelReviewSans"
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, windows_font))
        except Exception:
            font_name = "Helvetica"
    if font_name == "Helvetica":
        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            font_name = "STSong-Light"
        except Exception:
            pass

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=request.title,
        author="Yobboy File Server",
    )
    base = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=base["Title"], fontName=font_name, fontSize=27, leading=36, alignment=TA_CENTER, textColor=colors.HexColor("#202328"))
    heading = ParagraphStyle("SectionHeading", parent=base["Heading1"], fontName=font_name, fontSize=18, leading=24, textColor=colors.HexColor("#202328"))
    subheading = ParagraphStyle("SubHeading", parent=base["Heading2"], fontName=font_name, fontSize=13, leading=18, textColor=colors.HexColor("#2d8f67"), spaceBefore=0, spaceAfter=0, keepWithNext=False)
    body = ParagraphStyle("ReportBody", parent=base["BodyText"], fontName=font_name, fontSize=9.5, leading=14, textColor=colors.HexColor("#363b42"), alignment=TA_LEFT, spaceAfter=0)
    code = ParagraphStyle("ReportCode", parent=body, backColor=colors.HexColor("#f0f1f3"), borderPadding=5)
    label = ParagraphStyle("ViewLabel", parent=body, alignment=TA_CENTER, fontSize=10, leading=13)
    detail_label = ParagraphStyle("DetailLabel", parent=body, fontSize=9, leading=12, textColor=colors.HexColor("#737982"))
    streams: list[io.BytesIO] = []
    story: list[Any] = []

    story.extend([
        Spacer(1, 45 * mm),
        Paragraph(_pdf_mixed_text(request.title), title_style),
        Spacer(1, 10 * mm),
        Paragraph(f"模型文件：{_pdf_mixed_text(model_filename)}", ParagraphStyle("CoverMeta", parent=body, alignment=TA_CENTER, fontSize=12, leading=18)),
        Paragraph(f"生成时间：{_pdf_mixed_text(datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S'))}", ParagraphStyle("CoverTime", parent=body, alignment=TA_CENTER, fontSize=10, leading=16)),
        Spacer(1, 8 * mm),
        Paragraph(f"本报告包含 {request.preview_count} 张基础视图" + (f"和 {len(annotations)} 条评论审阅页。" if request.include_annotations else "。"), ParagraphStyle("CoverSummary", parent=body, alignment=TA_CENTER)),
        PageBreak(),
        Paragraph("第一部分  基础预览", heading),
        Paragraph("以下视图使用统一比例渲染完整装配体，不包含当前隐藏状态、爆炸位移或界面辅助元素。", body),
        Spacer(1, 5 * mm),
    ])

    view_items = list(request.views)
    for index, (view_id, view_label) in enumerate(view_items):
        if index > 0:
            story.extend([PageBreak(), Paragraph("第一部分  基础预览（续）", heading), Spacer(1, 3 * mm)])
        stream = io.BytesIO(preview_images[view_id])
        streams.append(stream)
        image = Image(stream, width=250 * mm, height=140.625 * mm)
        table = Table([[image], [Paragraph(view_label, label)]], colWidths=[255 * mm])
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d6d9dd")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("LEFTPADDING", (0, 0), (-1, -1), 2 * mm),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2 * mm),
            ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
        ]))
        story.append(table)

    if request.include_annotations:
        story.extend([
            PageBreak(),
            Paragraph("评论", heading),
        ])
        for index, annotation in enumerate(annotations, start=1):
            if index > 1:
                story.extend([PageBreak(), Paragraph("评论", heading)])
            image_bytes = annotation_images[annotation["id"]]
            if len(image_bytes) == 1:
                stream = io.BytesIO(image_bytes[0])
                streams.append(stream)
                image_grid: Any = Image(stream, width=210 * mm, height=118.125 * mm)
                image_grid.hAlign = "CENTER"
            else:
                image_cells = []
                for image_index in range(0, len(image_bytes), 2):
                    row = []
                    for data in image_bytes[image_index:image_index + 2]:
                        stream = io.BytesIO(data)
                        streams.append(stream)
                        row.append(Image(stream, width=105 * mm, height=59.0625 * mm))
                    if len(row) == 1:
                        row.append("")
                    image_cells.append(row)
                image_grid = Table(image_cells, colWidths=[107 * mm, 107 * mm], hAlign="CENTER")
                image_grid.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0.5 * mm),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 0.5 * mm),
                    ("TOPPADDING", (0, 0), (-1, -1), 0.5 * mm),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5 * mm),
                ]))
            markdown_flowables: list[Any] = [
                Paragraph("评论标题", detail_label),
                Paragraph(f"{index}. {html.escape(annotation['title'])}", subheading),
                Spacer(1, 1.5 * mm),
                Paragraph("评论内容", detail_label),
                Spacer(1, 0.8 * mm),
            ]
            has_markdown = False
            for kind, level, content in _markdown_blocks(annotation.get("markdown", "")):
                has_markdown = True
                if kind == "code":
                    markdown_flowables.append(Paragraph(html.escape(content).replace("\n", "<br/>"), code))
                else:
                    style = subheading if kind == "heading" else body
                    markdown_flowables.append(Paragraph(_inline_html(content), style))
                    markdown_flowables.append(Spacer(1, 1 * mm))
            if markdown_flowables and isinstance(markdown_flowables[-1], Spacer):
                markdown_flowables.pop()
            if not has_markdown:
                markdown_flowables.append(Paragraph("无正文", body))
            story.extend([Spacer(1, 1.5 * mm), image_grid, Spacer(1, 1 * mm), *markdown_flowables])

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor("#737982"))
        canvas.drawString(15 * mm, 8 * mm, request.title[:60])
        canvas.drawRightString(282 * mm, 8 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_docx_rfonts(element, latin_font: str, east_asia_font: str = "Microsoft YaHei") -> None:
    from docx.oxml.ns import qn

    fonts = element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin_font)
    fonts.set(qn("w:hAnsi"), latin_font)
    fonts.set(qn("w:eastAsia"), east_asia_font)
    fonts.set(qn("w:cs"), east_asia_font)
    for theme_name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{theme_name}"), None)


def _set_docx_style_font(style, latin_font: str = "Microsoft YaHei") -> None:
    style.font.name = latin_font
    _set_docx_rfonts(style.element, latin_font)


def _set_docx_run_font(run, latin_font: str = "Microsoft YaHei") -> None:
    run.font.name = latin_font
    _set_docx_rfonts(run._r, latin_font)


def _add_docx_inline(paragraph, children: list[Any]) -> None:
    bold = False
    italic = False
    for token in children:
        if isinstance(token, str):
            paragraph.add_run(token)
        elif token.type == "text":
            run = paragraph.add_run(token.content)
            run.bold = bold
            run.italic = italic
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type == "code_inline":
            run = paragraph.add_run(token.content)
            _set_docx_run_font(run, "Consolas")
        elif token.type == "image":
            paragraph.add_run(token.content or "[图片略]")


def _add_docx_markdown(container, markdown: str) -> None:
    from docx.shared import Pt

    for kind, level, content in _markdown_blocks(markdown):
        if kind == "code":
            paragraph = container.add_paragraph()
            run = paragraph.add_run(content.rstrip())
            _set_docx_run_font(run, "Consolas")
        else:
            style = f"Heading {min(level + 1, 6)}" if kind == "heading" else None
            paragraph = container.add_paragraph(style=style)
            _add_docx_inline(paragraph, content)
            if kind == "list":
                paragraph.paragraph_format.left_indent = 18
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1
    if not markdown.strip():
        paragraph = container.add_paragraph("无正文")
        paragraph.paragraph_format.space_after = Pt(0)


def build_docx(
    request: ModelReportRequest,
    model_filename: str,
    preview_images: dict[str, bytes],
    annotations: list[dict[str, Any]],
    annotation_images: dict[str, list[bytes]],
    glb_bytes: bytes | None,
    template_path: str | Path | None,
    separate_model_files: list[tuple[str, bytes, bytes, float]] | None = None,
) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Mm, Pt

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        _set_docx_style_font(document.styles[style_name])
    normal = document.styles["Normal"]
    normal.font.size = Pt(9.5)
    document.core_properties.title = request.title
    document.core_properties.author = "Yobboy File Server"

    cover = document.add_paragraph()
    cover.paragraph_format.space_before = Inches(1.55)
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = cover.add_run(request.title)
    _set_docx_run_font(title_run)
    title_run.bold = True
    title_run.font.size = Pt(27)
    meta = document.add_paragraph(f"模型文件：{model_filename}\n生成时间：{datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S')}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary = document.add_paragraph(
        f"本报告包含 {request.preview_count} 张基础视图"
        + (f"和 {len(annotations)} 条评论审阅页。" if request.include_annotations else "。")
    )
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("第一部分  基础预览", level=1)
    document.add_paragraph("以下视图使用统一比例渲染完整装配体，不包含当前隐藏状态、爆炸位移或界面辅助元素。")
    view_items = list(request.views)
    for index, (view_id, view_label) in enumerate(view_items):
        if index > 0:
            document.add_page_break()
            document.add_heading("第一部分  基础预览（续）", level=1)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(io.BytesIO(preview_images[view_id]), width=Inches(10.1))
        label = document.add_paragraph(view_label)
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if request.include_annotations:
        document.add_page_break()
        comment_heading = document.add_heading("评论", level=1)
        comment_heading.paragraph_format.space_before = Pt(0)
        comment_heading.paragraph_format.space_after = Pt(4)
        for index, annotation in enumerate(annotations, start=1):
            if index > 1:
                document.add_page_break()
                comment_heading = document.add_heading("评论", level=1)
                comment_heading.paragraph_format.space_before = Pt(0)
                comment_heading.paragraph_format.space_after = Pt(4)
            image_bytes = annotation_images[annotation["id"]]
            if len(image_bytes) == 1:
                image_paragraph = document.add_paragraph()
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_paragraph.paragraph_format.space_before = Pt(0)
                image_paragraph.paragraph_format.space_after = Pt(0)
                image_paragraph.add_run().add_picture(io.BytesIO(image_bytes[0]), width=Inches(8.4))
            else:
                image_table = document.add_table(rows=math.ceil(len(image_bytes) / 2), cols=2)
                image_table.autofit = False
                for image_index, data in enumerate(image_bytes):
                    image_item_cell = image_table.cell(image_index // 2, image_index % 2)
                    image_item_cell.width = Inches(4.25)
                    image_paragraph = image_item_cell.paragraphs[0]
                    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    image_paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(4.15))
            title_label = document.add_paragraph()
            title_label.paragraph_format.space_before = Pt(0)
            title_label.paragraph_format.space_after = Pt(0)
            title_run = title_label.add_run("评论标题")
            title_run.bold = True
            title_run.font.size = Pt(9)
            heading = document.add_paragraph()
            heading.style = document.styles["Heading 2"]
            heading.paragraph_format.space_before = Pt(0)
            heading.paragraph_format.space_after = Pt(2)
            heading.add_run(f"{index}. {annotation['title']}")
            content_label = document.add_paragraph()
            content_label.paragraph_format.space_before = Pt(2)
            content_label.paragraph_format.space_after = Pt(0)
            content_run = content_label.add_run("评论内容")
            content_run.bold = True
            content_run.font.size = Pt(9)
            _add_docx_markdown(document, annotation.get("markdown", ""))

    if request.include_3d:
        document.add_page_break()
        document.add_heading("第三部分  交互式 3D 模型", level=1)
        document.add_paragraph("请使用兼容的 Microsoft Word 桌面版打开本文档。拖动模型可旋转视角，滚轮可缩放；其他查看器将显示静态海报。")
        if request.separate_3d:
            document.add_heading("3.1 完整装配", level=2)
        placeholder = document.add_paragraph(PLACEHOLDER)
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for index, (name, _glb, _poster, _max_dimension) in enumerate(
            separate_model_files or (),
            start=2,
        ):
            document.add_page_break()
            document.add_heading(f"3.{index} {name}", level=2)
            document.add_paragraph("该模型从完整装配的顶层结构中单独导出，可在 Word 桌面版中旋转查看。")
            item_placeholder = document.add_paragraph(PLACEHOLDER)
            item_placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    document.save(buffer)
    if not request.include_3d:
        return buffer.getvalue()
    if glb_bytes is None:
        raise ModelReportValidationError("Word 导出缺少 GLB 模型")
    if template_path is None:
        raise ModelReportValidationError("3D Word 模板文件缺失")
    models = [(glb_bytes, preview_images["front"], request.model_max_dimension)]
    models.extend(
        (item_glb, poster, max_dimension)
        for _name, item_glb, poster, max_dimension in (separate_model_files or ())
    )
    return embed_models3d(
        buffer.getvalue(),
        template_path,
        models,
    )


def build_zip(files: dict[str, bytes], notes: str | None = None) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
        if notes:
            archive.writestr("导出说明.txt", notes.encode("utf-8"))
    return output.getvalue()


__all__ = [
    "Model3DEmbedError",
    "ModelReportRequest",
    "ModelReportValidationError",
    "build_docx",
    "build_pdf",
    "build_zip",
    "parse_report_request",
    "read_validated_glb",
    "read_validated_image",
    "safe_download_stem",
]
